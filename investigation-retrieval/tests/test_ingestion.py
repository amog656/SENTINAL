import json

import pymupdf
import pytest
from docx import Document

from app.ingestion.cleaner import clean_text_preserve_structure
from app.ingestion.parser import ParseError, parse_document, validate_file
from app.ingestion.service import ingest_document


def test_txt_parsing(tmp_path):
    file_path = tmp_path / "incident_notes.txt"
    file_path.write_text("orders-api latency increased\nv2.8.1", encoding="utf-8")

    text, document_type, file_size, page_count, metadata, extension = parse_document(str(file_path))

    assert document_type == "txt"
    assert file_size > 0
    assert page_count is None
    assert extension == ".txt"
    assert metadata["encoding_used"] == "utf-8"
    assert "orders-api latency" in text


def test_json_parsing_preserves_metadata(tmp_path):
    file_path = tmp_path / "incident.json"
    payload = {
        "document_id": "INC-1042",
        "service": "orders-api",
        "version": "v2.8.1",
        "description": "Latency increased after deployment.",
    }
    file_path.write_text(json.dumps(payload), encoding="utf-8")

    text, document_type, _, _, metadata, _ = parse_document(str(file_path))

    assert document_type == "json"
    assert "Latency increased" in text
    assert metadata["document_id"] == "INC-1042"
    assert metadata["service"] == "orders-api"
    assert metadata["original_json"] == payload


def test_pdf_parsing(tmp_path):
    file_path = tmp_path / "postmortem.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Payment service recovered after rollback.")
    doc.save(file_path)
    doc.close()

    text, document_type, _, page_count, _, _ = parse_document(str(file_path))

    assert document_type == "pdf"
    assert page_count == 1
    assert "PAGE 1" in text
    assert "Payment service" in text


def test_docx_parsing(tmp_path):
    file_path = tmp_path / "deployment_review.docx"
    doc = Document()
    doc.add_paragraph("Deployment DEP-882")
    doc.add_paragraph("orders-api v2.8.1 changed connection_pool.")
    doc.save(file_path)

    text, document_type, _, page_count, _, _ = parse_document(str(file_path))

    assert document_type == "docx"
    assert page_count is None
    assert "Deployment DEP-882" in text
    assert "connection_pool" in text


def test_text_cleaning_preserves_technical_identifiers():
    raw_text = "\r\n  orders-api   v2.8.1  \n\n\n  INC-1042\tDEP-882  \nconnection_pool  10:05 AM  "

    cleaned = clean_text_preserve_structure(raw_text)

    assert "orders-api v2.8.1" in cleaned
    assert "INC-1042 DEP-882" in cleaned
    assert "connection_pool 10:05 AM" in cleaned
    assert "\r" not in cleaned
    assert "\n\n\n\n" not in cleaned


def test_unsupported_extension(tmp_path):
    file_path = tmp_path / "notes.csv"
    file_path.write_text("unsupported", encoding="utf-8")

    with pytest.raises(ParseError, match="Unsupported file type"):
        validate_file(str(file_path))


def test_empty_file(tmp_path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("", encoding="utf-8")

    with pytest.raises(ParseError, match="Empty file"):
        validate_file(str(file_path))


def test_malformed_json(tmp_path):
    file_path = tmp_path / "bad.json"
    file_path.write_text("{not valid json", encoding="utf-8")

    with pytest.raises(ParseError, match="Invalid JSON"):
        parse_document(str(file_path))


def test_deterministic_document_id_for_same_content(tmp_path):
    first = tmp_path / "incident_notes.txt"
    first.write_text("Same incident evidence", encoding="utf-8")

    first_doc = ingest_document(str(first), filename="incident_notes.txt")
    second_doc = ingest_document(str(first), filename="incident_notes.txt")

    assert first_doc.document_id == second_doc.document_id
